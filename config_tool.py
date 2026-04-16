import json
import sys
import os
import re
import shutil
import csv
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

PROJECT_ROOT = Path(__file__).parent
DATA_DIR = PROJECT_ROOT / "data"
XLSX_PATH = PROJECT_ROOT / "xlsx" / "config.xlsx"
CSV_PATH = PROJECT_ROOT / "locale" / "translations.csv"

CONFIG_FILES = [
    "towers.json",
    "enemies.json",
    "events.json",
    "options.json",
    "stages.json",
    "traits.json",
    "endings.json",
    "economy_balance.json",
    "eras.json",
    "achievements.json",
    "family_backgrounds.json",
    "attributes.json",
    "game_config.json",
    "special_effects.json",
]

I18N_FIELDS = {
    "towers": ["tower_name", "tower_type", "tier"],
    "enemies": ["enemy_name", "enemy_type", "tier"],
    "events": ["event_name", "description"],
    "options": ["text"],
    "stages": ["stage_name", "description"],
    "traits": ["name", "description"],
    "endings": ["ending_name", "description", "ending_phrase"],
    "achievements": ["achievement_name", "description"],
    "attributes": ["display_name"],
    "family_backgrounds": ["family_name", "description", "traits.text"],
    "eras": ["era_name", "time_period"],
    "special_effects": ["effect_name", "description"],
}

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

JSON_ARRAY_PREFIX = "__json_arr__"
JSON_EMPTY_DICT = "__empty_dict__"
JSON_NULL = "__null__"

META_ROW_TAG = "__meta__"

I18N_COMMENT_TAG = "i18n_key:"


def load_translations():
    csv_map = {}
    zh_to_key = {}
    if not CSV_PATH.exists():
        print(f"  [!] 翻译文件不存在: {CSV_PATH}")
        return csv_map, zh_to_key
    with open(CSV_PATH, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        if not header:
            return csv_map, zh_to_key
        for row in reader:
            if len(row) < 2:
                continue
            key = row[0].strip()
            zh = row[1].strip() if len(row) > 1 else ""
            en = row[2].strip() if len(row) > 2 else ""
            csv_map[key] = {"zh": zh, "en": en}
            if zh:
                zh_to_key[zh] = key
    return csv_map, zh_to_key


def save_translations(csv_map):
    rows = [["keys", "zh", "en"]]
    for key in csv_map:
        entry = csv_map[key]
        rows.append([key, entry.get("zh", ""), entry.get("en", "")])
    with open(CSV_PATH, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def is_i18n_field(sheet_name, field_path):
    fields = I18N_FIELDS.get(sheet_name, [])
    for pattern in fields:
        if "." in pattern:
            if field_path.endswith(pattern) or pattern in field_path:
                return True
        else:
            if field_path == pattern or field_path.endswith("." + pattern):
                return True
    return False


def is_translation_key(value):
    if not isinstance(value, str):
        return False
    return bool(re.match(r'^[A-Z][A-Z0-9_]{2,}$', value))


def flatten(obj, prefix=""):
    if obj is None:
        return {prefix: JSON_NULL}
    if isinstance(obj, dict):
        if len(obj) == 0:
            return {prefix: JSON_EMPTY_DICT}
        result = {}
        for k, v in obj.items():
            new_key = f"{prefix}.{k}" if prefix else k
            if isinstance(v, (dict, list)):
                sub = flatten(v, new_key)
                result.update(sub)
            elif v is None:
                result[new_key] = JSON_NULL
            else:
                result[new_key] = v
        return result
    elif isinstance(obj, list):
        if len(obj) == 0:
            return {prefix: JSON_ARRAY_PREFIX + "[]"}
        first = obj[0]
        if isinstance(first, (dict, list)):
            result = {}
            for i, item in enumerate(obj):
                sub = flatten(item, f"{prefix}[{i}]")
                result.update(sub)
            return result
        else:
            return {prefix: JSON_ARRAY_PREFIX + json.dumps(obj, ensure_ascii=False)}
    else:
        return {prefix: obj}


def try_parse_value(value):
    if isinstance(value, str):
        if value == JSON_NULL:
            return None
        if value == JSON_EMPTY_DICT:
            return {}
        if value.startswith(JSON_ARRAY_PREFIX):
            json_str = value[len(JSON_ARRAY_PREFIX):]
            if json_str == "[]":
                return []
            try:
                return json.loads(json_str)
            except (json.JSONDecodeError, ValueError):
                return value
    return value


def unflatten(flat_dict):
    result = {}
    for compound_key, value in flat_dict.items():
        parts = parse_key_path(compound_key)
        parsed_value = try_parse_value(value)
        set_nested(result, parts, parsed_value)
    return result


def parse_key_path(key):
    parts = []
    current = ""
    i = 0
    while i < len(key):
        ch = key[i]
        if ch == ".":
            if current:
                parts.append(current)
                current = ""
        elif ch == "[":
            if current:
                parts.append(current)
                current = ""
            j = key.index("]", i)
            idx = key[i + 1 : j]
            parts.append(int(idx))
            i = j
        else:
            current += ch
        i += 1
    if current:
        parts.append(current)
    return parts


def set_nested(obj, parts, value):
    for i, part in enumerate(parts):
        is_last = i == len(parts) - 1
        if isinstance(part, int):
            while len(obj) <= part:
                obj.append(None)
            if is_last:
                obj[part] = value
            else:
                if obj[part] is None:
                    next_part = parts[i + 1]
                    obj[part] = [] if isinstance(next_part, int) else {}
                obj = obj[part]
        else:
            if is_last:
                obj[part] = value
            else:
                next_part = parts[i + 1]
                if part not in obj:
                    obj[part] = [] if isinstance(next_part, int) else {}
                obj = obj[part]


def detect_list_config(data):
    if not isinstance(data, dict):
        return None, data
    for key, value in data.items():
        if isinstance(value, list) and len(value) > 0 and isinstance(value[0], dict):
            return key, value
    return None, data


def sanitize_sheet_name(name):
    invalid_chars = ["\\", "/", "*", "?", ":", "[", "]"]
    for ch in invalid_chars:
        name = name.replace(ch, "_")
    return name[:31]


def write_sheet_from_json(ws, data, sheet_name="", csv_map=None):
    list_key, list_items = detect_list_config(data)

    if list_key and isinstance(list_items, list):
        version = data.get("version", "")
        meta_row = [META_ROW_TAG, f"version={version}", f"list_key={list_key}"]
        ws.append(meta_row)

        flat_rows = []
        all_keys = []
        for item in list_items:
            flat = flatten(item)
            flat_rows.append(flat)
            for k in flat:
                if k not in all_keys:
                    all_keys.append(k)

        ws.append(all_keys)
        for row_data in flat_rows:
            row = []
            for k in all_keys:
                val = row_data.get(k, "")
                if csv_map and is_i18n_field(sheet_name, k) and is_translation_key(val):
                    zh_text = csv_map.get(val, {}).get("zh", val)
                    row.append(zh_text)
                else:
                    row.append(val)
            ws.append(row)

        if csv_map:
            for col_idx, k in enumerate(all_keys, 1):
                if not is_i18n_field(sheet_name, k):
                    continue
                for row_idx in range(3, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    flat_row = flat_rows[row_idx - 3]
                    original_val = flat_row.get(k, "")
                    if is_translation_key(original_val):
                        cell.comment = Comment(f"{I18N_COMMENT_TAG}{original_val}", "config_tool")
    else:
        flat = flatten(data)
        ws.append(["key", "value"])
        for k in sorted(flat.keys()):
            val = flat[k]
            if csv_map and is_i18n_field(sheet_name, k) and is_translation_key(val):
                zh_text = csv_map.get(val, {}).get("zh", val)
                ws.append([k, zh_text])
            else:
                ws.append([k, val])

    for col in ws.columns:
        max_length = 0
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
            cell.border = THIN_BORDER
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 60)

    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN

        if list_key:
            for cell in ws[2]:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = HEADER_ALIGN


def read_sheet_to_json(ws, sheet_name="", csv_map=None, zh_to_key=None):
    rows = list(ws.iter_rows(values_only=False))
    if not rows:
        return None, []

    first_row_vals = [c.value for c in rows[0]]
    is_list_config = False
    list_key = None
    version = ""
    csv_updates = []

    if first_row_vals and first_row_vals[0] == META_ROW_TAG:
        is_list_config = True
        for cell in first_row_vals[1:]:
            if cell and str(cell).startswith("version="):
                version = str(cell).split("=", 1)[1]
            elif cell and str(cell).startswith("list_key="):
                list_key = str(cell).split("=", 1)[1]

    if is_list_config and list_key:
        if len(rows) < 3:
            return None, csv_updates
        headers = [str(c.value) for c in rows[1] if c.value is not None]
        items = []
        for row in rows[2:]:
            row_vals = [c.value for c in row]
            if row_vals[0] is None:
                continue
            row_dict = {}
            for h, cell, val in zip(headers, row, row_vals):
                if val is None or val == "":
                    continue
                if is_i18n_field(sheet_name, h):
                    resolved = _resolve_i18n_value(cell, val, csv_map, zh_to_key, sheet_name, h)
                    if isinstance(resolved, tuple):
                        row_dict[h] = resolved[0]
                        if resolved[1]:
                            csv_updates.append(resolved[1])
                    else:
                        row_dict[h] = resolved
                else:
                    row_dict[h] = val
            unflat = unflatten(row_dict)
            items.append(unflat)

        result = {}
        if version:
            result["version"] = version
        result[list_key] = items
        return result, csv_updates
    else:
        flat = {}
        for row in rows:
            row_vals = [c.value for c in row]
            key_val = row_vals[0]
            if key_val and key_val != "key":
                val = row_vals[1] if len(row_vals) > 1 else ""
                if is_i18n_field(sheet_name, str(key_val)):
                    cell = row[1] if len(row) > 1 else None
                    resolved = _resolve_i18n_value(cell, val, csv_map, zh_to_key, sheet_name, str(key_val))
                    if isinstance(resolved, tuple):
                        flat[str(key_val)] = resolved[0]
                        if resolved[1]:
                            csv_updates.append(resolved[1])
                    else:
                        flat[str(key_val)] = resolved
                else:
                    flat[str(key_val)] = val
            elif key_val == "key":
                continue
        return unflatten(flat), csv_updates


def _resolve_i18n_value(cell, value, csv_map, zh_to_key, sheet_name, field_path):
    if not isinstance(value, str) or not value.strip():
        return value, None

    text = str(value).strip()

    if is_translation_key(text):
        return text, None

    comment_key = _get_comment_key(cell)
    if comment_key:
        if csv_map and comment_key in csv_map:
            original_zh = csv_map[comment_key].get("zh", "")
            if text != original_zh and original_zh:
                csv_map[comment_key]["zh"] = text
                en_val = csv_map[comment_key].get("en", "")
                if en_val and not en_val.startswith("[NEEDS UPDATE]"):
                    csv_map[comment_key]["en"] = f"[NEEDS UPDATE] {en_val}"
                return comment_key, ("update", comment_key, text)
            return comment_key, None
        return comment_key, None

    if zh_to_key and text in zh_to_key:
        return zh_to_key[text], None

    if _is_likely_enum_value(text):
        return text, None

    new_key = _generate_i18n_key(sheet_name, field_path, text)
    if csv_map is not None:
        csv_map[new_key] = {"zh": text, "en": ""}
    print(f"  [!] 新翻译键: {new_key} = {text}")
    return new_key, ("new", new_key, text)


def _get_comment_key(cell):
    if cell is None or cell.comment is None:
        return None
    comment_text = cell.comment.text
    tag = I18N_COMMENT_TAG
    if tag in comment_text:
        for line in comment_text.split("\n"):
            line = line.strip()
            if line.startswith(tag):
                key = line[len(tag):].strip()
                if is_translation_key(key):
                    return key
    return None


def _is_likely_enum_value(text):
    if re.match(r'^[a-z][a-z0-9_]*$', text):
        return True
    return False


def _generate_i18n_key(sheet_name, field_path, text):
    prefix_map = {
        "towers": "TOWER",
        "enemies": "ENEMY",
        "events": "EVENT",
        "options": "OPTION",
        "stages": "STAGE",
        "traits": "TRAIT",
        "endings": "ENDING",
        "achievements": "ACHIEVEMENT",
        "attributes": "ATTR",
        "family_backgrounds": "FAMILY",
        "eras": "ERA",
        "special_effects": "SE",
    }
    prefix = prefix_map.get(sheet_name, sheet_name.upper())
    field_suffix = field_path.split(".")[-1].upper()
    suffix_map = {
        "NAME": "NAME",
        "TOWER_NAME": "NAME",
        "ENEMY_NAME": "NAME",
        "EVENT_NAME": "NAME",
        "STAGE_NAME": "NAME",
        "ENDING_NAME": "NAME",
        "ACHIEVEMENT_NAME": "NAME",
        "FAMILY_NAME": "NAME",
        "ERA_NAME": "NAME",
        "EFFECT_NAME": "NAME",
        "DESCRIPTION": "DESC",
        "TEXT": "TEXT",
        "ENDING_PHRASE": "PHRASE",
        "DISPLAY_NAME": "DISPLAY_NAME",
        "TIME_PERIOD": "TIME_PERIOD",
        "TOWER_TYPE": "TYPE",
        "ENEMY_TYPE": "TYPE",
        "TIER": "TIER",
    }
    key_suffix = suffix_map.get(field_suffix, field_suffix)
    hash_part = format(abs(hash(text)), "X")[:6]
    return f"{prefix}_NEW_{hash_part}_{key_suffix}"


def validate_sheet(ws, original_data):
    errors = []
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        errors.append("sheet 为空")
        return errors

    first_row = rows[0]
    is_list_config = False
    list_key = None

    if first_row and first_row[0] == META_ROW_TAG:
        is_list_config = True
        for cell in first_row[1:]:
            if cell and str(cell).startswith("list_key="):
                list_key = str(cell).split("=", 1)[1]

    if is_list_config and list_key:
        if len(rows) < 3:
            errors.append("缺少数据行")
            return errors

        headers = [str(h) for h in rows[1] if h is not None]
        id_column = f"{list_key[:-1]}_id" if list_key.endswith("s") else None

        if not id_column or id_column not in headers:
            for candidate in headers:
                if candidate.endswith("_id"):
                    id_column = candidate
                    break

        original_ids = set()
        if list_key in original_data and isinstance(original_data[list_key], list):
            for item in original_data[list_key]:
                if isinstance(item, dict):
                    if id_column and id_column in item:
                        original_ids.add(str(item[id_column]))
                    else:
                        for k in item:
                            if k.endswith("_id"):
                                original_ids.add(str(item[k]))
                                break

        imported_ids = set()
        for row in rows[2:]:
            if row[0] is None:
                continue
            row_dict = dict(zip(headers, row))
            if id_column and id_column in row_dict:
                imported_ids.add(str(row_dict[id_column]))

        if original_ids and id_column:
            missing = original_ids - imported_ids
            if missing:
                errors.append(
                    f"缺少以下 {id_column}: {', '.join(sorted(missing))}。"
                    f"导入会删除这些条目，请确认是否故意删除"
                )

        for i, row in enumerate(rows[2:], start=3):
            if row[0] is None:
                continue
            row_dict = dict(zip(headers, row))
            for k, v in row_dict.items():
                if v is None or v == "":
                    continue
                if k.endswith("_id") and not isinstance(v, str):
                    errors.append(f"行 {i}: {k} 应为字符串，实际为 {type(v).__name__}: {v}")

    return errors


def cmd_export():
    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    csv_map, _ = load_translations()
    wb = Workbook()
    wb.remove(wb.active)

    print("导出 JSON -> config.xlsx (i18n 模式):")
    for filename in CONFIG_FILES:
        json_path = DATA_DIR / filename
        if not json_path.exists():
            print(f"  !! {filename} 不存在，跳过")
            continue

        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        sheet_name = sanitize_sheet_name(filename.replace(".json", ""))
        ws = wb.create_sheet(title=sheet_name)

        try:
            write_sheet_from_json(ws, data, sheet_name=sheet_name, csv_map=csv_map)
            i18n_count = len(I18N_FIELDS.get(sheet_name, []))
            if i18n_count > 0:
                print(f"  -> sheet [{sheet_name}] (含 {i18n_count} 个 i18n 字段)")
            else:
                print(f"  -> sheet [{sheet_name}]")
        except Exception as e:
            print(f"  !! {filename} 导出失败: {e}")
            wb.remove(ws)

    wb.save(XLSX_PATH)
    print(f"导出完成! 文件: {XLSX_PATH}")


def cmd_import():
    if not XLSX_PATH.exists():
        print(f"{XLSX_PATH.name} 不存在，请先运行 export")
        return

    csv_map, zh_to_key = load_translations()
    wb = load_workbook(XLSX_PATH)
    print("导入 config.xlsx -> JSON (i18n 模式):")

    success_count = 0
    fail_count = 0
    all_csv_updates = []

    sheet_to_file = {}
    for filename in CONFIG_FILES:
        sheet_name = sanitize_sheet_name(filename.replace(".json", ""))
        sheet_to_file[sheet_name] = filename

    for ws in wb.worksheets:
        sheet_name = ws.title
        filename = sheet_to_file.get(sheet_name, f"{sheet_name}.json")
        json_path = DATA_DIR / filename

        if not json_path.exists():
            print(f"  !! {filename} 不存在，跳过 sheet [{sheet_name}]")
            fail_count += 1
            continue

        with open(json_path, "r", encoding="utf-8") as f:
            original_data = json.load(f)

        errors = validate_sheet(ws, original_data)
        if errors:
            print(f"  !! sheet [{sheet_name}] 验证失败:")
            for e in errors:
                print(f"     - {e}")
            print(f"  !! 跳过 sheet [{sheet_name}]，原配置未修改")
            fail_count += 1
            continue

        result, csv_updates = read_sheet_to_json(ws, sheet_name=sheet_name, csv_map=csv_map, zh_to_key=zh_to_key)
        if result is None:
            print(f"  !! sheet [{sheet_name}] 读取失败，跳过")
            fail_count += 1
            continue

        all_csv_updates.extend(csv_updates)

        if isinstance(original_data, dict) and isinstance(result, dict):
            for k in original_data:
                if k == "version":
                    continue
                if k not in result:
                    result[k] = original_data[k]

        bak_path = json_path.with_suffix(".json.bak")
        shutil.copy2(json_path, bak_path)

        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent="\t")

        i18n_count = len(I18N_FIELDS.get(sheet_name, []))
        update_info = f" ({len(csv_updates)} 处翻译更新)" if csv_updates else ""
        print(f"  -> {filename} (备份: {bak_path.name}){update_info}")
        success_count += 1

    wb.close()

    if all_csv_updates:
        save_translations(csv_map)
        update_count = sum(1 for u in all_csv_updates if u[0] == "update")
        new_count = sum(1 for u in all_csv_updates if u[0] == "new")
        parts = []
        if update_count:
            parts.append(f"{update_count} 处中文已更新")
        if new_count:
            parts.append(f"{new_count} 个新翻译键已追加")
        print(f"  [CSV] translations.csv 已同步: {', '.join(parts)}")

    print(f"导入完成! 成功: {success_count}, 失败: {fail_count}")


def main():
    if len(sys.argv) < 2:
        print("用法:")
        print("  python config_tool.py export   - 导出 data/*.json -> config.xlsx (翻译键→中文)")
        print("  python config_tool.py import   - 导入 config.xlsx -> data/*.json (中文→翻译键)")
        sys.exit(1)

    cmd = sys.argv[1].lower()
    if cmd == "export":
        cmd_export()
    elif cmd == "import":
        cmd_import()
    else:
        print(f"未知命令: {cmd}")
        print("可用命令: export, import")
        sys.exit(1)


if __name__ == "__main__":
    main()
